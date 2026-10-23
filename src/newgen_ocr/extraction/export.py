"""JSON, CSV and Excel export of extraction results (docs/EXTRACTION.md).

Spreadsheet formula injection is neutralised: in CSV a text cell starting with ``= + - @``,
tab or CR is prefixed with ``'`` (plain numbers such as ``-250.00`` are left as they are);
in Excel every text cell is written as a string, never a formula. Missing values export as
empty cells, never as 0.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter

from newgen_ocr.extraction.models import Evidence, ExtractionResult, FieldResult, TableResult

_NUMBER = re.compile(r"^-?\d+(\.\d+)?$")
_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(_DANGEROUS) and not _NUMBER.match(text):
        return "'" + text
    return text


def _evidence(evidence: list[Evidence]) -> str:
    parts = []
    for e in evidence:
        box = f"{e.bbox.x0:.0f},{e.bbox.y0:.0f},{e.bbox.x1:.0f},{e.bbox.y1:.0f}{e.units}"
        where = f"p{e.page_number} {e.source.value} [{box}]"
        if e.table_id is not None:
            where += f" table {e.table_id} r{e.row}c{e.col}"
        parts.append(where)
    return "; ".join(parts)


FIELD_COLUMNS = [
    "field", "type", "required", "status", "value", "raw", "method", "confidence",
    "pages", "evidence", "notes", "errors",
]  # fmt: skip


def _field_row(f: FieldResult) -> list[object]:
    return [
        f.name, f.type, f.required, f.status.value, f.value, f.raw,
        f.method.value if f.method else None, f.confidence,
        ",".join(str(p) for p in sorted({e.page_number for e in f.evidence})),
        _evidence(f.evidence), "; ".join(f.notes), "; ".join(f.errors),
    ]  # fmt: skip


def _table_rows(table: TableResult) -> tuple[list[str], list[list[object]]]:
    header = ["record", "page", "row", *table.columns, "statuses", "evidence"]
    rows: list[list[object]] = []
    for rec in table.records:
        values = [rec.values[c].value if c in rec.values else None for c in table.columns]
        statuses = "; ".join(
            f"{c}={rec.values[c].status.value}"
            for c in table.columns
            if c in rec.values and rec.values[c].status.value != "extracted"
        )
        evidence = _evidence(
            [v.evidence for v in rec.values.values() if v.evidence is not None][:1]
        )
        rows.append([rec.index, rec.page_number, rec.row, *values, statuses, evidence])
    return header, rows


def to_json(result: ExtractionResult) -> bytes:
    return result.model_dump_json(indent=2).encode()


def _csv(header: list[str], rows: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([safe_text(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")  # BOM: Excel opens UTF-8 (Indic text) correctly


def fields_csv(result: ExtractionResult) -> bytes:
    return _csv(FIELD_COLUMNS, [_field_row(f) for f in result.fields])


def table_csv(result: ExtractionResult, name: str) -> bytes | None:
    table = next((t for t in result.tables if t.name == name), None)
    if table is None:
        return None
    return _csv(*_table_rows(table))


def _typed(value: object, kind: str | None) -> object:
    """Excel-native values: amounts/numbers as numbers, dates as dates, others as text."""
    if value is None or value == "":
        return None
    if kind in ("amount", "number", "percent") and isinstance(value, str) and _NUMBER.match(value):
        return Decimal(value)
    if kind == "date" and isinstance(value, str):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _put(cell: Cell, value: object) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # a string, never a formula
    elif isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"


def _sheet(book: Workbook, title: str, header: list[str], rows: list[list[object]]) -> None:
    sheet = book.create_sheet(title[:31])
    for c, name in enumerate(header, start=1):
        _put(sheet.cell(row=1, column=c), name)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            _put(sheet.cell(row=r, column=c), value)
    for c in range(1, len(header) + 1):
        sheet.column_dimensions[get_column_letter(c)].width = 18
    sheet.freeze_panes = "A2"


def to_xlsx(result: ExtractionResult) -> bytes:
    book = Workbook()
    book.remove(book.active)
    summary: list[list[object]] = [
        ["document_id", result.document_id],
        ["document_type", result.schema_id],
        ["schema_version", result.schema_version],
        ["classification", result.classification.status.value],
        ["status", result.status.value],
        ["needs_review", result.review.needs_review],
        ["review_reasons", "; ".join(r.value for r in result.review.reasons)],
        ["valid", result.validation.valid],
        ["layout_run", result.layout_run],
        ["ocr_run", result.ocr_run],
        ["created_at", result.created_at.isoformat()],
        ["pipeline_version", result.pipeline_version],
    ]
    _sheet(book, "Summary", ["item", "value"], summary)
    types = {f.name: f.type for f in result.fields}
    rows = [_field_row(f) for f in result.fields]
    for row in rows:
        row[4] = _typed(row[4], types.get(str(row[0])))
    _sheet(book, "Fields", FIELD_COLUMNS, rows)
    for table in result.tables:
        header, table_rows = _table_rows(table)
        for row in table_rows:
            for i, name in enumerate(table.columns, start=3):
                row[i] = _typed(row[i], table.column_types.get(name))
        _sheet(book, f"Table {table.name}", header, table_rows)
    _sheet(
        book,
        "Checks",
        ["check", "kind", "status", "message", "expected", "actual", "difference", "rows"],
        [
            [
                c.id,
                c.kind,
                c.status.value,
                c.message,
                c.expected,
                c.actual,
                c.difference,
                ",".join(str(r) for r in c.rows),
            ]
            for c in result.checks
        ],
    )
    _sheet(
        book,
        "Validation",
        ["code", "location", "message"],
        [
            [i.code, "/".join(str(p) for p in i.location), i.message]
            for i in result.validation.issues
        ],
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
