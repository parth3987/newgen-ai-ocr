"""Extraction service (stored runs, review, masking, native PDF text) and exports."""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from newgen_ocr.audit import AuditLog
from newgen_ocr.config import Settings
from newgen_ocr.errors import AppError, ErrorCode
from newgen_ocr.extraction import export
from newgen_ocr.extraction.extract import Options
from newgen_ocr.extraction.models import (
    ClassificationStatus,
    EvidenceSource,
    ExtractionResult,
    ExtractionReviewReason,
    ExtractionStatus,
)
from newgen_ocr.extraction.schema import SchemaRegistry
from newgen_ocr.extraction.service import ExtractionService
from newgen_ocr.imaging.runtime import runtime_info
from newgen_ocr.ingestion.service import IngestionService, UploadInfo
from newgen_ocr.layout.models import LayoutRun, PageLayout, PageLayoutStatus
from newgen_ocr.ocr.models import OcrPage, OcrRun, PageOcrStatus
from newgen_ocr.storage.local import DocumentStore
from tests.fixtures.builders import PageSpec, build_pdf
from tests.fixtures.extraction import INVOICE, ITEMS, store_document, table

TODAY = date(2024, 6, 1)


@pytest.fixture
def store(settings: Settings) -> DocumentStore:
    return DocumentStore(settings.data_dir)


def _service(store: DocumentStore, settings: Settings, **options: bool) -> ExtractionService:
    return ExtractionService(
        store, AuditLog(settings.data_dir), SchemaRegistry(), Options(**options)
    )


def test_complete_invoice_is_stored_and_audited(store: DocumentStore, settings: Settings) -> None:
    doc = store_document(settings, store, INVOICE, [table(ITEMS)])
    result = _service(store, settings).extract(doc, "auto", today=TODAY)
    assert result.status is ExtractionStatus.COMPLETE, result.review
    assert (result.schema_id, result.schema_version) == ("gst_invoice", "1.0")
    assert result.classification.status is ClassificationStatus.CLASSIFIED
    assert result.validation.valid and not result.review.needs_review
    assert result.record["grand_total"] == "1180.00"
    assert result.tables[0].records[1].values["description"].value == "Toner"
    assert store.list_extractions(doc) == ["auto"]
    assert store.read_extraction(doc, "auto").record == result.record
    event = json.loads((settings.data_dir / "audit" / "audit.jsonl").read_text().splitlines()[-1])
    assert (event["event"], event["document_type"], event["status"]) == (
        "extraction.completed",
        "gst_invoice",
        "complete",
    )
    assert "Grand Total" not in json.dumps(event)  # no document content in the audit log


def test_review_reasons(store: DocumentStore, settings: Settings) -> None:
    blocks = [["TAX INVOICE"], ["Invoice Date: 14/03/2024", "Grand Total: 1,190.00"]]
    doc = store_document(settings, store, blocks)
    result = _service(store, settings).extract(doc, "auto", today=TODAY)
    assert result.status is ExtractionStatus.NEEDS_REVIEW
    reasons = set(result.review.reasons)
    assert ExtractionReviewReason.REQUIRED_FIELD_MISSING in reasons
    assert ExtractionReviewReason.TABLE_NOT_FOUND in reasons
    assert "invoice_number" in result.review.details
    assert result.record["invoice_number"] is None


def test_unclassified_and_unknown_type(store: DocumentStore, settings: Settings) -> None:
    doc = store_document(settings, store, [["Dear customer, the office is closed on Monday."]])
    result = _service(store, settings).extract(doc, "auto")
    assert result.status is ExtractionStatus.UNCLASSIFIED
    assert result.fields == [] and result.schema_id is None
    assert result.review.reasons == [ExtractionReviewReason.DOCUMENT_TYPE_UNKNOWN]
    with pytest.raises(AppError) as exc:
        _service(store, settings).extract(doc, "auto", "passport")
    assert exc.value.code is ErrorCode.UNKNOWN_DOCUMENT_TYPE
    with pytest.raises(AppError) as exc:
        _service(store, settings).extract(doc, "other")
    assert exc.value.code is ErrorCode.LAYOUT_RUN_NOT_FOUND
    with pytest.raises(AppError) as exc:
        store.read_extraction(doc, "other")
    assert exc.value.code is ErrorCode.EXTRACTION_NOT_FOUND


def test_requested_type_and_low_confidence(store: DocumentStore, settings: Settings) -> None:
    doc = store_document(settings, store, [["Receipt No: R-1", "Date: 05/04/2024", "Total: 50.00"]])
    service = ExtractionService(
        store, AuditLog(settings.data_dir), SchemaRegistry(), min_evidence_confidence=0.99
    )
    result = service.extract(doc, "auto", "receipt", today=TODAY)
    assert result.classification.status is ClassificationStatus.REQUESTED
    assert ExtractionReviewReason.LOW_CONFIDENCE_EVIDENCE in result.review.reasons


AADHAAR = [
    ["Government of India", "Unique Identification Authority of India"],
    ["Name: Anita Desai", "DOB: 02/11/1985", "Aadhaar No: 2468 1357 9026"],
]


def test_aadhaar_is_masked_everywhere(store: DocumentStore, settings: Settings) -> None:
    doc = store_document(settings, store, AADHAAR)
    result = _service(store, settings).extract(doc, "auto", today=TODAY)
    number = next(f for f in result.fields if f.name == "aadhaar_number")
    assert number.value == "XXXX XXXX 9026"
    stored = (settings.data_dir / "derived" / doc / "extraction" / "auto.json").read_text()
    assert "2468 1357" not in stored and "246813579026" not in stored
    assert "2468" not in export.fields_csv(result).decode()
    full = _service(store, settings, mask_aadhaar=False).extract(doc, "auto", today=TODAY)
    assert next(f for f in full.fields if f.name == "aadhaar_number").value == "246813579026"


def test_native_pdf_text_is_used(store: DocumentStore, settings: Settings) -> None:
    text = "TAX INVOICE\nInvoice No: INV-9\nInvoice Date: 01/02/2024\nGrand Total: 2,360.00"
    manifest = IngestionService(settings, store, AuditLog(settings.data_dir)).ingest(
        io.BytesIO(build_pdf([PageSpec(text)])), UploadInfo("n.pdf", None)
    )
    doc = manifest.document_id
    now = datetime.now(tz=UTC)
    skipped = {"status": PageOcrStatus.SKIPPED, "skip_reason": "native_text"}
    store.write_ocr_run(
        doc,
        "auto",
        OcrRun(
            document_id=doc,
            requested_engine="auto",
            engines=[],
            languages=["en"],
            preprocessing_profile=None,
            created_at=now,
            pipeline_version="t",
            runtime=runtime_info(),
            total_ms=0,
            pages=[OcrPage(page_number=1, **skipped)],
        ),
    )
    store.write_layout_run(
        doc,
        LayoutRun(
            document_id=doc,
            ocr_run="auto",
            models=[],
            parameters={},
            created_at=now,
            pipeline_version="t",
            runtime=runtime_info(),
            total_ms=0,
            pages=[
                PageLayout(
                    page_number=1, status=PageLayoutStatus.SKIPPED, skip_reason="native_text"
                )
            ],
        ),
    )
    result = _service(store, settings).extract(doc, "auto", "gst_invoice", today=TODAY)
    fields = {f.name: f for f in result.fields}
    assert fields["invoice_number"].value == "INV-9"
    assert fields["grand_total"].value == "2360.00"
    evidence = fields["grand_total"].evidence[0]
    assert (evidence.source, evidence.units) == (EvidenceSource.PDF_TEXT, "pt")
    assert "embedded PDF text used" in result.source_notes[0]


def _result(store: DocumentStore, settings: Settings) -> ExtractionResult:
    rows = [*ITEMS, ['=HYPERLINK("http://x")', "4802", "1", "-5.00", "-5.00"]]
    doc = store_document(settings, store, INVOICE, [table(rows)])
    return _service(store, settings).extract(doc, "auto", today=TODAY)


def test_json_and_csv_exports(store: DocumentStore, settings: Settings) -> None:
    result = _result(store, settings)
    assert ExtractionResult.model_validate_json(export.to_json(result)) == result
    data = export.fields_csv(result)
    assert data.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == export.FIELD_COLUMNS
    by_name = {r[0]: r for r in rows[1:]}
    assert by_name["grand_total"][4] == "1180.00"
    assert by_name["due_date"][3:5] == ["missing", ""]  # missing stays empty, not 0
    items = list(
        csv.reader(io.StringIO(export.table_csv(result, "line_items").decode("utf-8-sig")))
    )
    formula = items[3]
    assert formula[3].startswith("'=")  # formula injection neutralised
    assert formula[6] == formula[8] == "-5.00"  # negative numbers are left alone
    assert export.table_csv(result, "nope") is None
    assert export.safe_text("@cmd") == "'@cmd" and export.safe_text("-3") == "-3"


def test_xlsx_export(store: DocumentStore, settings: Settings) -> None:
    result = _result(store, settings)
    book = load_workbook(io.BytesIO(export.to_xlsx(result)))
    assert book.sheetnames == ["Summary", "Fields", "Table line_items", "Checks", "Validation"]
    fields = {row[0].value: row for row in book["Fields"].iter_rows(min_row=2)}
    assert float(fields["grand_total"][4].value) == 1180.0
    assert fields["invoice_date"][4].value.date() == date(2024, 3, 14)
    assert fields["due_date"][4].value is None
    items = list(book["Table line_items"].iter_rows(min_row=2, values_only=True))
    assert items[2][3] == '=HYPERLINK("http://x")'
    cell = book["Table line_items"]["D4"]
    assert cell.data_type == "s"  # stored as text, never as a formula
