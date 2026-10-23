"""Structured extraction and exports through the HTTP API."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from newgen_ocr.api.app import create_app
from newgen_ocr.config import Settings
from newgen_ocr.evaluation.client import authorized
from newgen_ocr.evaluation.extraction_corpus import gst_invoice
from newgen_ocr.storage.local import DocumentStore
from tests.conftest import MODELS_DIR, MODELS_LOCK
from tests.fixtures.extraction import INVOICE, ITEMS, store_document, table

DOCS = "/api/v1/documents"


@pytest.fixture
def api_settings(tmp_path: Path) -> Settings:
    data = tmp_path / "data"
    data.mkdir()
    return Settings(  # type: ignore[call-arg]
        environment="test",
        data_dir=data,
        models_dir=MODELS_DIR,
        models_lock=MODELS_LOCK,
        _env_file=None,
    )


@pytest.fixture
def client(api_settings: Settings) -> Iterator[TestClient]:
    with TestClient(create_app(api_settings), raise_server_exceptions=False) as c:
        yield authorized(c)


def test_schemas(client: TestClient) -> None:
    listed = {s["id"]: s for s in client.get("/api/v1/extraction/schemas").json()}
    assert set(listed) == {"gst_invoice", "receipt", "bank_statement", "pan_card", "aadhaar_card"}
    assert "grand_total" in listed["gst_invoice"]["fields"]
    assert listed["bank_statement"]["tables"] == ["transactions"]
    schema = client.get("/api/v1/extraction/schemas/pan_card").json()
    assert schema["fields"][0]["name"] == "pan"
    missing = client.get("/api/v1/extraction/schemas/passport")
    assert (missing.status_code, missing.json()["error"]["code"]) == (422, "UNKNOWN_DOCUMENT_TYPE")


def test_extraction_and_exports(client: TestClient, api_settings: Settings) -> None:
    store = DocumentStore(api_settings.data_dir)
    doc = store_document(api_settings, store, INVOICE, [table(ITEMS)])
    response = client.post(f"{DOCS}/{doc}/extraction", json={"layout_run": "auto"})
    assert response.status_code == 201, response.text
    body = response.json()
    assert (body["schema_id"], body["status"]) == ("gst_invoice", "complete")
    assert client.get(f"{DOCS}/{doc}/extraction").json() == ["auto"]
    assert client.get(f"{DOCS}/{doc}/extraction/auto").json()["record"] == body["record"]

    base = f"{DOCS}/{doc}/extraction/auto/export"
    as_json = client.get(base)
    assert as_json.headers["content-type"] == "application/json"
    assert as_json.json()["record"]["grand_total"] == "1180.00"
    fields = client.get(base, params={"format": "csv"})
    assert fields.headers["content-disposition"].endswith('-auto.csv"')
    rows = list(csv.reader(io.StringIO(fields.content.decode("utf-8-sig"))))
    assert rows[0][0] == "field"
    items = client.get(base, params={"format": "csv", "table": "line_items"})
    assert "Toner" in items.content.decode("utf-8-sig")
    unknown = client.get(base, params={"format": "csv", "table": "nope"})
    assert unknown.json()["error"]["code"] == "UNSUPPORTED_EXPORT"
    xlsx = client.get(base, params={"format": "xlsx"})
    assert load_workbook(io.BytesIO(xlsx.content)).sheetnames[0] == "Summary"
    assert client.get(base, params={"format": "pdf"}).status_code == 422

    forced = client.post(f"{DOCS}/{doc}/extraction", json={"document_type": "receipt"})
    assert forced.json()["classification"]["status"] == "requested"
    bad = client.post(f"{DOCS}/{doc}/extraction", json={"document_type": "passport"})
    assert bad.json()["error"]["code"] == "UNKNOWN_DOCUMENT_TYPE"
    no_layout = client.post(f"{DOCS}/{doc}/extraction", json={"layout_run": "tesseract"})
    assert no_layout.json()["error"]["code"] == "LAYOUT_RUN_NOT_FOUND"
    gone = client.get(f"{DOCS}/{doc}/extraction/tesseract")
    assert gone.json()["error"]["code"] == "EXTRACTION_NOT_FOUND"


@pytest.mark.models
def test_live_invoice_end_to_end(client: TestClient) -> None:
    image, truth = gst_invoice("intra", 200)
    buf = io.BytesIO()
    Image.fromarray(image).save(buf, "PNG")
    doc = client.post(DOCS, files={"file": ("inv.png", buf.getvalue(), "image/png")}).json()[
        "document_id"
    ]
    assert client.post(f"{DOCS}/{doc}/ocr", json={"languages": ["en"]}).status_code == 201
    assert client.post(f"{DOCS}/{doc}/layout", json={}).status_code == 201
    result = client.post(f"{DOCS}/{doc}/extraction", json={}).json()
    assert result["schema_id"] == "gst_invoice"
    fields = {f["name"]: f for f in result["fields"]}
    for name, value in truth.fields.items():
        if value is None:
            assert fields[name]["status"] != "extracted", name  # never invented
        else:
            assert fields[name]["value"] == value, name
    assert {c["id"]: c["status"] for c in result["checks"]}["grand_total"] == "passed"
    evidence = fields["grand_total"]["evidence"][0]
    assert evidence["source"] == "ocr_line" and evidence["line_indices"]
